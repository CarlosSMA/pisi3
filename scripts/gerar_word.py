from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt


def clean_cell(value: str) -> str:
    return re.sub(r"[`*_]", "", value.strip())


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [
        [clean_cell(cell) for cell in line.strip().strip("|").split("|")]
        for line in lines
    ]
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        table.rows[0].cells[index].text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if index < len(cells):
                cells[index].text = value


def add_markdown(document: Document, content: str) -> None:
    lines = content.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                document.add_paragraph("\n".join(code_lines), style="No Spacing")
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_heading(heading.group(2), level=level)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        elif line.strip():
            document.add_paragraph(line)
        index += 1


def build_document(source_dir: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    document.add_heading("Dicionário de dados epidemiológicos", level=0)
    document.add_paragraph(
        "Documento gerado a partir das análises dos arquivos CSV do projeto."
    )
    documents = sorted(source_dir.glob("*.md"))
    for position, source in enumerate(documents):
        if position:
            document.add_section(WD_SECTION.NEW_PAGE)
        add_markdown(document, source.read_text(encoding="utf-8"))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Documentos Markdown incorporados: {len(documents)}")
    print(f"Documento Word gerado em: {output}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Gera o dicionário de dados em Word.")
    parser.add_argument("--source-dir", type=Path, default=Path("docs"))
    parser.add_argument(
        "--output", type=Path, default=Path("docs") / "dicionario_dados.docx"
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_document(args.source_dir, args.output)


if __name__ == "__main__":
    main()